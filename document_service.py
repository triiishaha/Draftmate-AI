
from io import BytesIO
from pathlib import Path

from pypdf import PdfReader
from docx import Document


def extract_document(filename: str, file_bytes: bytes) -> str:
    """
    Extract text from a PDF or DOCX file.
    Returns the extracted text as a string.
    """

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        pdf = PdfReader(BytesIO(file_bytes))

        if pdf.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported.")

        pages = []

        for page in pdf.pages:
            page_text = page.extract_text() or ""
            pages.append(page_text)

        text = "\n\n".join(pages)

    elif extension == ".docx":
        document = Document(BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        tables = []

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                tables.append(" | ".join(cells))

        text = "\n".join(paragraphs + tables)

    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    if not text.strip():
        raise ValueError(
            "No text found. The file may be scanned and require OCR."
        )

    return text

