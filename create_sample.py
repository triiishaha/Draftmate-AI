
from docx import Document

doc = Document()
doc.add_heading("SERVICE AGREEMENT", level=1)
doc.add_heading("1. Payment Terms", level=2)
doc.add_paragraph("The client must pay within 30 days.")
doc.add_heading("1.1 Late Fees", level=3)
doc.add_paragraph("Late payments may incur additional fees.")
doc.add_heading("2. Termination", level=2)
doc.add_paragraph("Either party may terminate this agreement.")

doc.save("sample_agreement.docx")
print("Sample DOCX created!")
