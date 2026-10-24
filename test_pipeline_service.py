
from pathlib import Path
from docx import Document

from pipeline_service import process_document_command

# Create a test agreement
doc = Document()
doc.add_heading("SERVICE AGREEMENT", level=1)
doc.add_heading("1. Payment Terms", level=2)
doc.add_paragraph("The client must pay within 30 days.")
doc.add_heading("1.1 Late Fees", level=3)
doc.add_paragraph("Late payments may incur additional fees.")

doc.save("pipeline_test.docx")

# Read the document bytes
file_path = Path("pipeline_test.docx")
file_bytes = file_path.read_bytes()

# Process a user command
result = process_document_command(
    filename=file_path.name,
    file_bytes=file_bytes,
    transcript="Rewrite clause 1.1 in simpler language"
)

print("Action:", result["action"])
print("Target clause:", result["target_clause"])
print("Context found:", result["context"]["found"])
print("Selected text:", result["context"]["text"])
print("Estimated tokens:", result["estimated_tokens"])

assert result["action"] == "rewrite_clause"
assert result["context"]["found"] is True

print("\nPipeline service test passed!")
