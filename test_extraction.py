
from io import BytesIO
from docx import Document

from document_service import extract_document


def create_test_docx():
    document = Document()
    document.add_heading("SERVICE AGREEMENT", level=1)
    document.add_heading("1. Payment Terms", level=2)
    document.add_paragraph(
        "The client must pay within 30 days."
    )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


result = extract_document("sample.docx", create_test_docx())

print("Extraction successful!")
print(result)
