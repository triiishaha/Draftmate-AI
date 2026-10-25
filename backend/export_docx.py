"""Export a draft as a real Word (.docx) file: POST /export/docx {text, title}."""
from __future__ import annotations

import io
import re

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

router = APIRouter(prefix="/export", tags=["export"])
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ExportIn(BaseModel):
    text: str
    title: str = "DraftMate document"


def _is_heading(line: str) -> bool:
    t = line.strip()
    if not t or len(t) > 80 or t.endswith((".", ",", ";")):
        return False
    return t.isupper() or t.startswith("#") or bool(re.match(r"^(\d+[.)]|Article|Clause|Schedule)\s+[A-Z]", t) and len(t.split()) <= 7)


def build_docx(text: str, title: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.3
    doc.core_properties.title = title[:120]
    doc.core_properties.author = "DraftMate AI"

    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n").strip()):
        lines = block.split("\n")
        if len(lines) == 1 and _is_heading(lines[0]):
            p = doc.add_paragraph()
            r = p.add_run(lines[0].strip().lstrip("#").strip())
            r.bold = True
            p.paragraph_format.keep_with_next = True
            if lines[0].strip().isupper():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        p = doc.add_paragraph()
        for i, line in enumerate(lines):
            run = p.add_run(line.strip())
            run.bold = bool(re.match(r"^\*\*.+\*\*$", line.strip()))
            if run.bold:
                run.text = line.strip().strip("*")
            if i < len(lines) - 1:
                run.add_break()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.post("/docx")
def export_docx(body: ExportIn):
    if not body.text.strip():
        raise HTTPException(422, "There is nothing to export yet.")
    if len(body.text) > 500_000:
        raise HTTPException(413, "The draft is too long to export.")
    try:
        data = build_docx(body.text, body.title)
    except ImportError:
        raise HTTPException(501, "Word export needs python-docx. Run: python3 -m pip install python-docx")
    name = re.sub(r"[^A-Za-z0-9._-]+", "-", body.title).strip("-") or "DraftMate"
    return Response(data, media_type=DOCX, headers={"Content-Disposition": f'attachment; filename="{name}.docx"'})
