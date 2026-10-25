import io
import pytest
import doc_convert
from doc_convert import convert_bytes, ConvertError


def make_docx():
    import docx
    d = docx.Document()
    d.add_heading("Services Agreement", 1)
    d.add_paragraph("This agreement is dated 1 November 2026.")
    d.add_heading("1. Payment", 2)
    d.add_paragraph("The Client shall pay Rs. 50,000 within 15 days.")
    d.add_paragraph("Late fee applies", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Amount", "Fee", "50,000"
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


def make_pdf(text="Termination. Either party may end this agreement with 30 days notice.", pages=2):
    from reportlab.pdfgen import canvas
    b = io.BytesIO(); c = canvas.Canvas(b)
    for i in range(pages):
        c.drawString(72, 750, f"{text} (page {i+1})"); c.showPage()
    c.save()
    return b.getvalue()


def blank_pdf():
    from reportlab.pdfgen import canvas
    b = io.BytesIO(); c = canvas.Canvas(b); c.rect(50, 50, 100, 100); c.showPage(); c.save()
    return b.getvalue()


@pytest.mark.parametrize("prefer", [False, True])
def test_docx_keeps_headings_lists_tables(prefer):
    r = convert_bytes("a.docx", make_docx(), prefer_markitdown=prefer)
    md = r["markdown"]
    assert "# Services Agreement" in md and "Payment" in md and "Rs. 50,000" in md
    assert "Item" in md and "Amount" in md and "|" in md
    assert r["tokens_estimate"] > 0 and r["engine"] in {"built-in", "MarkItDown"}


@pytest.mark.parametrize("prefer", [False, True])
def test_pdf_text_extracted(prefer):
    r = convert_bytes("a.pdf", make_pdf(), prefer_markitdown=prefer)
    assert "30 days notice" in r["markdown"] and r["pages"] == 2 and r["warnings"] == []


def test_scanned_style_pdf_is_flagged_not_guessed():
    r = convert_bytes("scan.pdf", blank_pdf(), prefer_markitdown=False)
    assert r["markdown"] == "" and any("OCR" in w for w in r["warnings"])


def test_plain_text_html_csv():
    assert convert_bytes("a.txt", "Hello\r\n\r\n\r\n\r\nWorld".encode())["markdown"] == "Hello\n\nWorld"
    h = convert_bytes("a.html", b"<html><head><style>x{}</style></head><body><h1>Title</h1><p>Para</p><script>bad()</script><ul><li>One</li></ul></body></html>", prefer_markitdown=False)["markdown"]
    assert "# Title" in h and "Para" in h and "- One" in h and "bad()" not in h
    c = convert_bytes("a.csv", b"name,fee\nA,10\nB,20", prefer_markitdown=False)["markdown"]
    assert c.splitlines()[0] == "| name | fee |" and "| A | 10 |" in c


def test_bad_inputs_are_refused_clearly():
    with pytest.raises(ConvertError, match="not supported"):
        convert_bytes("virus.exe", b"MZ")
    with pytest.raises(ConvertError, match="empty"):
        convert_bytes("a.txt", b"")
    with pytest.raises(ConvertError, match="too large"):
        convert_bytes("a.txt", b"x" * (doc_convert.MAX_BYTES + 1))
    with pytest.raises(ConvertError):
        convert_bytes("broken.pdf", b"not a pdf", prefer_markitdown=False)
    with pytest.raises(ConvertError):
        convert_bytes("broken.docx", b"not a docx", prefer_markitdown=False)


def test_path_in_filename_is_dropped():
    assert convert_bytes("../../etc/x.txt", b"hi")["filename"] == "x.txt"


def test_pptx_without_markitdown_gives_install_hint(monkeypatch):
    monkeypatch.setattr(doc_convert, "_markitdown", lambda e, d: None)
    with pytest.raises(ConvertError, match="MarkItDown"):
        convert_bytes("a.pptx", b"PK")


def test_upload_route():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    app = FastAPI(); app.include_router(doc_convert.router); c = TestClient(app)
    r = c.post("/doc/convert", files={"file": ("contract.docx", make_docx())})
    assert r.status_code == 200 and "Services Agreement" in r.json()["markdown"]
    assert c.post("/doc/convert", files={"file": ("x.exe", b"MZ")}).status_code == 422
