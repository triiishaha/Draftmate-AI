import io, os, sys, types
os.environ.setdefault("GEMINI_API_KEY", "test-key")
try:
    from google import genai  # noqa: F401
except ImportError:
    g, gg = types.ModuleType("google"), types.ModuleType("google.genai")
    gg.Client = lambda **k: object(); g.genai = gg
    sys.modules["google"], sys.modules["google.genai"] = g, gg
from docx import Document
from fastapi.testclient import TestClient
from main import app

c = TestClient(app)
TEXT = "LEGAL NOTICE\n\nTo,\nMr. Rao\n\n1. DEMAND\n\nYou are required to return the deposit within 15 days."

def test_docx_is_valid_word_file():
    r = c.post("/export/docx", json={"text": TEXT, "title": "Legal Notice"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "attachment" in r.headers["content-disposition"] and r.headers["content-disposition"].endswith('.docx"')
    paras = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert "LEGAL NOTICE" in paras and any("return the deposit" in p for p in paras)
    heading = [p for p in Document(io.BytesIO(r.content)).paragraphs if p.text == "LEGAL NOTICE"][0]
    assert heading.runs[0].bold

def test_empty_rejected():
    assert c.post("/export/docx", json={"text": "  "}).status_code == 422


def test_markdown_becomes_formatting_not_symbols():
    md = "### Summary\n\n**Parties:** Aditi and Khushi\n\n* First point\n* Second point\n\n---\n\nPlain *emphasis* here."
    r = c.post("/export/docx", json={"text": md, "title": "T"})
    d = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "###" not in text and "**" not in text and "---" not in text and "*First" not in text
    assert "Summary" in text and "First point" in text
    bold_runs = [run.text for p in d.paragraphs for run in p.runs if run.bold]
    assert "Parties:" in bold_runs


def test_pdf_is_valid_and_markdown_safe():
    md = "### Summary\n\n**Parties:** Aditi and Khushi\n\n* First point\n* Second point\n\n---\n\nPlain text."
    r = c.post("/export/pdf", json={"text": md, "title": "Legal Notice"})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert r.headers["content-disposition"].endswith('.pdf"')
    assert c.post("/export/pdf", json={"text": " "}).status_code == 422
